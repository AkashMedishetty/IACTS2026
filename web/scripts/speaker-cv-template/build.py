"""
IACTS TechnoCollege CME 2026 — speaker profile (CV) template.

    python scripts/speaker-cv-template/build.py <out_dir>
    (needs python-pptx and Pillow)

Two 16:9 slides, one per amount of information the team has:

  1. Full profile  — photo, name and role on a crimson panel; profile,
                     qualifications, experience, interests, awards; their session.
  2. Brief profile — for when a speaker sends little: photo in the crimson
                     circle of the event artwork, name, role, institution and
                     their session. Looks complete with just those.

Both carry the event header. Every field holds sample text in its final style,
so typing over it keeps the look; boxes shrink text that runs long. The photo is
a stand-in: right-click › Change Picture. Duplicate a slide per speaker.
"""
import io
import sys
from pathlib import Path

from PIL import Image, ImageDraw, ImageFilter
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

ROOT = Path(__file__).resolve().parents[2]  # web/
LOGO = ROOT / "public/brand/iacts-technocollege-2026-logo.png"

CRIMSON = RGBColor(0xB3, 0x12, 0x2A)
DEEP = RGBColor(0x8B, 0x00, 0x00)
WINE = RGBColor(0x5F, 0x07, 0x17)
GOLD = RGBColor(0xA9, 0x79, 0x3A)
GOLD_SOFT = RGBColor(0xE3, 0xD3, 0xB5)
IVORY = RGBColor(0xFB, 0xF7, 0xF1)
BLUSH = RGBColor(0xF8, 0xE9, 0xED)
INK = RGBColor(0x17, 0x12, 0x0F)
SOFT = RGBColor(0x6A, 0x54, 0x5A)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

DISPLAY = "Impact"   # condensed display face close to the event's Anton; on every Mac and Windows
BODY = "Calibri"

W, H = 13.333, 7.5
HEAD_H = 1.0


# ── primitives ────────────────────────────────────────────────────────────────

def shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=1.0):
    s = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        s.fill.background()
    elif not isinstance(fill, RGBColor):  # (top, bottom) pair: vertical gradient, light at the top
        s.fill.gradient()
        s.fill.gradient_angle = 90
        stops = s.fill.gradient_stops
        stops[0].color.rgb, stops[0].position = fill[1], 0.0
        stops[1].color.rgb, stops[1].position = fill[0], 1.0
    else:
        s.fill.solid()
        s.fill.fore_color.rgb = fill
    if line is None:
        s.line.fill.background()
    else:
        s.line.color.rgb = line
        s.line.width = Pt(line_w)
    s.shadow.inherit = False
    if s.has_text_frame:
        s.text_frame.text = ""
    return s


def rule(slide, x1, y, x2, color=GOLD, width=1.0):
    """A hairline as a filled sliver: connectors pick up the theme's line effects."""
    h = width / 72
    return shape(slide, MSO_SHAPE.RECTANGLE, x1, y - h / 2, x2 - x1, h, fill=color)


def text(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, after=0, shrink=False):
    """paras: list of paragraphs; each a list of runs (text, font, size, color, bold[, tracking pt])."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    tf.margin_left = tf.margin_right = Inches(0.04)
    tf.margin_top = tf.margin_bottom = Inches(0.02)
    for i, runs in enumerate(paras):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(after)
        for t, font, size, color, bold, *track in runs:
            r = p.add_run()
            r.text = t
            r.font.name = font
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.bold = bold
            if track:
                r._r.get_or_add_rPr().set("spc", str(int(track[0] * 100)))
    if shrink:  # PowerPoint shrinks the text when someone types more than fits
        tf.auto_size = MSO_AUTO_SIZE.TEXT_TO_FIT_SHAPE
    return tb


def bullets(tb, color=GOLD):
    for p in tb.text_frame.paragraphs:
        pPr = p._p.get_or_add_pPr()
        pPr.set("marL", str(Inches(0.2)))
        pPr.set("indent", str(-Inches(0.2)))
        clr = pPr.makeelement(qn("a:buClr"), {})
        clr.append(clr.makeelement(qn("a:srgbClr"), {"val": str(color)}))
        pPr.append(clr)
        pPr.append(pPr.makeelement(qn("a:buFont"), {"typeface": "Arial"}))
        pPr.append(pPr.makeelement(qn("a:buChar"), {"char": "•"}))


def dots(slide, x, y, color=GOLD, n=3, step=0.17, d=0.05):
    """The small dot grid from the event artwork."""
    for i in range(n):
        for j in range(n):
            shape(slide, MSO_SHAPE.OVAL, x + i * step, y + j * step, d, d, fill=color)


def photo_png(bg=(0xEF, 0xE4, 0xD8), tone=(0xD9, 0xC5, 0xAE)):
    """Square stand-in portrait (a soft silhouette) for the team to swap out."""
    n = 900
    im = Image.new("RGB", (n, n), bg)
    d = ImageDraw.Draw(im)
    d.ellipse((n * .33, n * .2, n * .67, n * .56), fill=tone)
    d.ellipse((n * .13, n * .62, n * .87, n * 1.3), fill=tone)
    im = im.filter(ImageFilter.GaussianBlur(1.2))
    buf = io.BytesIO()
    im.save(buf, "PNG")
    buf.seek(0)
    return buf


def round_photo(slide, cx, cy, d):
    """Circular photo frame centred on (cx, cy): stays round after Change Picture."""
    pic = slide.shapes.add_picture(photo_png(), Inches(cx - d / 2), Inches(cy - d / 2), Inches(d), Inches(d))
    pic.auto_shape_type = MSO_SHAPE.OVAL
    pic.name = "Speaker photo"
    return pic


def ring(slide, cx, cy, d, color, width):
    return shape(slide, MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, line=color, line_w=width)


# ── shared header ─────────────────────────────────────────────────────────────

def header(slide, label):
    shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, W, HEAD_H, fill=IVORY)
    slide.shapes.add_picture(str(LOGO), Inches(0.42), Inches(0.13), height=Inches(HEAD_H - 0.26))
    text(slide, 6.0, 0.16, W - 6.42, 0.48, [[(label, DISPLAY, 24, CRIMSON, False, 3)]],
         align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    text(slide, 6.0, 0.6, W - 6.42, 0.3, [[
        ("IACTS TechnoCollege CME 2026", BODY, 11, INK, True),
        ("   ·   23 – 25 October 2026   ·   Hyderabad", BODY, 11, SOFT, False),
    ]], align=PP_ALIGN.RIGHT, anchor=MSO_ANCHOR.MIDDLE)
    band = shape(slide, MSO_SHAPE.RECTANGLE, 0, HEAD_H, W, 0.07, fill=CRIMSON)
    band.fill.gradient()
    band.fill.gradient_angle = 0
    st = band.fill.gradient_stops
    st[0].color.rgb, st[0].position = DEEP, 0.0
    st[1].color.rgb, st[1].position = GOLD, 1.0


def session_card(slide, x, y, w, h, title_size=16):
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, w, h, fill=BLUSH)
    shape(slide, MSO_SHAPE.RECTANGLE, x, y, 0.08, h, fill=CRIMSON)
    text(slide, x + 0.28, y + 0.12, w - 0.45, h - 0.2, [
        [("SPEAKING AT IACTS TECHNOCOLLEGE CME 2026", BODY, 10.5, CRIMSON, True, 1.5)],
        [("Title of the talk", BODY, title_size, INK, True)],
        [("Day · Date  |  Time  |  Hall", BODY, 12, SOFT, False)],
    ], after=3, shrink=True)


def notes(slide, which):
    slide.notes_slide.notes_text_frame.text = (
        f"{which}\n\n"
        "How to fill this slide\n"
        "• Photo: right-click the photo › Change Picture › From a File. It stays round. "
        "If the face is off-centre: Picture Format › Crop, then drag the photo inside the circle.\n"
        "• Text: click any text and type over it — font, size and colour stay. Long text shrinks to fit.\n"
        "• Not enough information? Delete the sections you cannot fill, or use the Brief profile slide.\n"
        "• One speaker per slide: right-click the slide thumbnail › Duplicate Slide, then fill the copy."
    )


# ── slide 1: full profile ─────────────────────────────────────────────────────

def section(slide, x, y, w, h, title, paras, bulleted=False):
    shape(slide, MSO_SHAPE.RECTANGLE, x, y + 0.09, 0.1, 0.1, fill=CRIMSON)
    text(slide, x + 0.18, y, w - 0.18, 0.28, [[(title.upper(), BODY, 11.5, CRIMSON, True, 1.5)]],
         anchor=MSO_ANCHOR.MIDDLE)
    rule(slide, x, y + 0.34, x + w, color=GOLD_SOFT, width=0.75)
    tb = text(slide, x, y + 0.42, w, h - 0.42, [[(t, BODY, 13, INK, False)] for t in paras], after=2, shrink=True)
    if bulleted:
        bullets(tb)
    return tb


def full_profile(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, W, H, fill=WHITE)
    header(s, "SPEAKER PROFILE")

    # Crimson identity panel
    px, pw = 0, 4.7
    top = HEAD_H + 0.07
    shape(s, MSO_SHAPE.RECTANGLE, px, top, pw, H - top, fill=(CRIMSON, WINE))
    dots(s, pw - 0.72, top + 0.25, color=GOLD_SOFT)
    cx, cy, d = pw / 2, top + 1.78, 2.75
    ring(s, cx, cy, d + 0.55, GOLD_SOFT, 0.75)
    ring(s, cx, cy, d + 0.2, GOLD, 3)
    round_photo(s, cx, cy, d)

    text(s, 0.3, cy + d / 2 + 0.28, pw - 0.6, 0.62, [[("DR. FULL NAME", DISPLAY, 28, WHITE, False, 1)]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, shrink=True)
    rule(s, cx - 0.45, cy + d / 2 + 1.0, cx + 0.45, color=GOLD, width=2)
    text(s, 0.3, cy + d / 2 + 1.12, pw - 0.6, H - (cy + d / 2 + 1.12) - 0.25, [
        [("Designation", BODY, 16, GOLD_SOFT, True)],
        [("Department", BODY, 13, WHITE, False)],
        [("Institution", BODY, 14.5, WHITE, True)],
        [("City, State", BODY, 12.5, GOLD_SOFT, False)],
    ], align=PP_ALIGN.CENTER, after=2, shrink=True)

    # CV sections
    x, cw = 5.2, W - 5.2 - 0.5
    y = top + 0.3
    section(s, x, y, cw, 1.15, "Profile", [
        "A short introduction in two or three lines — the speaker's work, focus and the programmes they lead.",
    ])
    section(s, x, y + 1.25, cw, 0.8, "Qualifications", ["MBBS, MS, MCh (CTVS) / DNB — Institution"])
    section(s, x, y + 2.15, cw, 1.3, "Experience", [
        "Current position — Institution (Year – present)",
        "Previous position — Institution (Years)",
        "__ years in cardiothoracic and vascular surgery",
    ], bulleted=True)
    half = (cw - 0.4) / 2
    section(s, x, y + 3.55, half, 1.05, "Areas of Interest", ["Speciality and procedures of focus"])
    section(s, x + half + 0.4, y + 3.55, half, 1.05, "Awards & Memberships", ["Awards, honours and societies"])

    session_card(s, x, H - 1.42, cw, 1.12)
    notes(s, "FULL PROFILE — for speakers who have sent a CV.")


# ── slide 2: brief profile ────────────────────────────────────────────────────

def brief_profile(prs):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    shape(s, MSO_SHAPE.RECTANGLE, 0, 0, W, H, fill=IVORY)
    header(s, "SPEAKER PROFILE")

    # The crimson circle of the event artwork, holding the photo
    cx, cy, d = 9.55, 4.3, 5.3
    ring(s, cx, cy, d + 0.75, GOLD, 1.25)
    shape(s, MSO_SHAPE.OVAL, cx - d / 2, cy - d / 2, d, d, fill=(CRIMSON, WINE))
    dots(s, 12.35, 1.45)
    ring(s, cx, cy, 3.95, WHITE, 4)
    round_photo(s, cx, cy, 3.85)

    x, w = 0.85, 5.5
    text(s, x, 1.75, w, 0.35, [[("OUR SPEAKER", BODY, 13, GOLD, True, 4)]])
    text(s, x, 2.1, w, 1.45, [[("DR. FULL NAME", DISPLAY, 44, CRIMSON, False, 1)]],
         anchor=MSO_ANCHOR.BOTTOM, shrink=True)
    rule(s, x + 0.05, 3.72, x + 1.25, color=GOLD, width=2.5)
    text(s, x, 3.88, w, 1.45, [
        [("Designation", BODY, 19, INK, True)],
        [("Institution", BODY, 16, DEEP, True)],
        [("City, State", BODY, 14, SOFT, False)],
    ], after=4, shrink=True)

    session_card(s, x, 5.55, w, 1.3, title_size=17)
    notes(s, "BRIEF PROFILE — for speakers who have sent only a photo, name and position.")


def build(out: Path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(W), Inches(H)
    cp = prs.core_properties
    cp.title = "IACTS TechnoCollege CME 2026 — Speaker Profile Template"
    cp.author = cp.last_modified_by = "IACTS TechnoCollege CME 2026"
    cp.subject = "IACTS TechnoCollege CME 2026 — Hyderabad, 23–25 October 2026"
    cp.category = "Speaker CV template"
    cp.comments = ""
    full_profile(prs)
    brief_profile(prs)
    prs.save(out)


if __name__ == "__main__":
    out_dir = Path(sys.argv[1] if len(sys.argv) > 1 else ".").expanduser()
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / "IACTS_TechnoCollege_2026_Speaker_CV_Template.pptx"
    build(out)
    print("wrote", out)
