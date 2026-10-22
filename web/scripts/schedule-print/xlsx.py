"""
The printable scientific schedule as an Excel workbook — same table, colours
and wording as the PDF (build.mjs), read from the rows build.mjs writes.

    node scripts/schedule-print/build.mjs
    python scripts/schedule-print/xlsx.py "<output.xlsx>" [schedule|workshop]     (needs openpyxl, Pillow)

Sheet 1 is the scientific schedule (both days), sheet 2 the Breakthrough
Sessions. Both are set up to print: A4 landscape, one page wide, the column
headings repeated on every page, page breaks where the PDF has them.
"""
import json
import sys
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.pagebreak import Break

HERE = Path(__file__).resolve().parent
ROOT = HERE.parents[1]  # web/
LOGO = ROOT / "public/brand/iacts-technocollege-2026-logo.png"
# Which document: the scientific schedule (default) or the workshop day.
DOC = sys.argv[2] if len(sys.argv) > 2 else "schedule"
DATA = json.loads((HERE / f"out/{DOC}.json").read_text())
TITLE = DATA.get("title", "SCIENTIFIC SCHEDULE")
HEADINGS = ["Timings", "S. No.", "Name", "Session Topic"]

DEEP, CRIMSON, WINE, GOLD = "8B0000", "B3122A", "5F0717", "A9793A"
INK, SOFT = "17120F", "6A545A"
F_INFO, F_SES, F_ALT, F_BRK, LINE = "F6EBDA", "F8E9ED", "FFFCF7", "EAD8B8", "D9C29C"
FONT = "Arial"

side = Side(style="thin", color=LINE)
BOX = Border(left=side, right=side, top=side, bottom=side)
fill = lambda c: PatternFill("solid", start_color=c, end_color=c)
font = lambda size=10, bold=False, color=INK: Font(name=FONT, size=size, bold=bold, color=color)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT = Alignment(horizontal="left", vertical="center", wrap_text=True, indent=1)
rich = lambda *parts: CellRichText(*[TextBlock(InlineFont(rFont=FONT, sz=s, b=b, color=c), t) for t, s, b, c in parts])

WIDTHS = {"A": 20, "B": 8, "C": 36, "D": 92}       # characters
LINE_PT = 17                                        # one wrapped line at 10 pt


def lines(text, width):
    """Rough wrapped-line count, to size rows (Excel does not grow merged rows)."""
    return sum(max(1, -(-len(part) // max(1, int(width * 1.05)))) for part in str(text).split("\n"))


class Sheet:
    def __init__(self, ws):
        self.ws, self.r = ws, 1

    def cell(self, col, value, f=None, fnt=None, al=CENTER):
        c = self.ws[f"{col}{self.r}"]
        c.value = value
        if fnt: c.font = fnt
        c.alignment = al
        if f: c.fill = fill(f)
        return c

    def span(self, a, b, value, f=None, fnt=None, al=CENTER):
        for col in "ABCD"["ABCD".index(a):"ABCD".index(b) + 1]:
            c = self.ws[f"{col}{self.r}"]
            c.border = BOX
            if f: c.fill = fill(f)
        if a != b:
            self.ws.merge_cells(f"{a}{self.r}:{b}{self.r}")
        return self.cell(a, value, f, fnt, al)

    def boxes(self, f=None):
        for col in "ABCD":
            c = self.ws[f"{col}{self.r}"]
            c.border = BOX
            if f: c.fill = fill(f)

    def height(self, n_lines, pad=12):
        self.ws.row_dimensions[self.r].height = n_lines * LINE_PT + pad

    def next(self):
        self.r += 1


def header(sh, widths):
    """Logo at the left, theme at the right — above the table, not repeated."""
    ws = sh.ws
    img = XLImage(str(LOGO))
    img.width, img.height = round(img.width * 62 / img.height), 62   # 62 px tall, true proportions
    ws.add_image(img, "A1")
    last = list(widths)[-1]
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 22
    ws.row_dimensions[3].height = 8
    c = ws[f"{last}1"]
    c.value = rich(("THE FUTURE ", 20, True, CRIMSON), ("IS NOW", 20, True, GOLD))
    c.alignment = Alignment(horizontal="right", vertical="bottom")
    c = ws[f"{last}2"]
    c.value = "23 – 25 OCTOBER 2026  ·  HYDERABAD"
    c.font = font(9, True, SOFT)
    c.alignment = Alignment(horizontal="right", vertical="top")
    for col in widths:
        ws[f"{col}3"].border = Border(bottom=Side(style="medium", color=CRIMSON))
    sh.r = 4


def schedule(ws):
    sh = Sheet(ws)
    for col, w in WIDTHS.items():
        ws.column_dimensions[col].width = w
    header(sh, WIDTHS)

    # Column headings + title (the headings repeat on every printed page)
    head_row = sh.r
    for col, t in zip("ABCD", HEADINGS):
        sh.cell(col, t, DEEP, font(10, True, "FFFFFF")).border = BOX
    sh.height(1, 10); sh.next()
    sh.span("A", "D", TITLE, CRIMSON, font(14, True, "FFFFFF")); sh.height(1, 14); sh.next()

    current_day, alt = None, False
    for i, page in enumerate(DATA["pages"]):
        day = page["day"]
        if i:  # page breaks where the PDF has them, so no session is split
            ws.row_breaks.append(Break(id=sh.r - 1))
        if day["title"] != current_day:
            if current_day is not None:
                sh.span("A", "D", TITLE, CRIMSON, font(14, True, "FFFFFF")); sh.height(1, 14); sh.next()
            current_day = day["title"]
            sh.span("A", "D", f'{day["title"]} - {day["date"]}', F_INFO, font(11.5, True, WINE)); sh.height(1, 10); sh.next()
            sh.span("A", "D", f'Venue : {day["venue"]}', F_INFO, font(11.5, True, WINE)); sh.height(1, 10); sh.next()
        for r in page["rows"]:
            k = r["k"]
            if k == "session":
                if r.get("chairs"):
                    sh.span("A", "D", rich((r["title"] + "\n", 12, True, CRIMSON), ("Chairpersons - " + r["chairs"], 10, False, INK)), F_SES)
                    sh.height(1 + lines("Chairpersons - " + r["chairs"], sum(WIDTHS.values())), 10)
                else:  # a hall heading on the workshop day
                    sh.span("A", "D", r["title"], F_SES, font(12.5, True, CRIMSON)); sh.height(1, 12)
                alt = False
            elif k == "talk":
                bg = F_ALT if alt else "FFFFFF"
                alt = not alt
                sh.boxes(bg)
                sh.cell("A", r["time"], bg, font(10), LEFT)
                sh.cell("B", r["n"], bg, font(10))
                names = "\n".join(r["name"]) if isinstance(r["name"], list) else r["name"]
                sh.cell("C", names, bg, font(10, True), LEFT)
                topic = rich((r["badge"] + ": ", 10, True, INK), (r["topic"], 10, False, INK)) if r.get("badge") else r["topic"]
                sh.cell("D", topic, bg, font(10))
                sh.height(max(lines(r["topic"], WIDTHS["D"]), lines(names, WIDTHS["C"])))
            elif k in ("break", "gala"):
                sh.boxes(F_BRK)
                sh.cell("A", r["time"], F_BRK, font(10, True), LEFT)
                sh.span("C", "D", r["label"], F_BRK, font(11.5, True, CRIMSON if k == "gala" else WINE))
                sh.height(1, 10)
            elif k == "special":  # Quiz heading, then the Quiz Master row
                sh.boxes(F_BRK)
                sh.cell("A", r["time"], F_BRK, font(10, True), LEFT)
                sh.span("C", "D", r["label"], F_BRK, font(11.5, True, WINE)); sh.height(1, 10); sh.next()
                bg = F_ALT if alt else "FFFFFF"
                sh.boxes(bg)
                sh.cell("B", r["n"], bg, font(10)); sh.cell("C", r["name"], bg, font(10, True), LEFT)
                sh.cell("D", r["topic"], bg, font(10)); sh.height(1)
            elif k == "ceremony":
                sh.boxes(F_SES)
                sh.cell("A", r["time"], F_SES, font(10, True), LEFT)
                g = r["guests"]
                rows3 = ["        ".join(g[i:i + 3]) for i in range(0, len(g), 3)]
                sh.cell("D", rich((r["title"] + "\n", 10.5, True, CRIMSON), ("\n".join(rows3), 10, True, INK)), F_SES)
                sh.height(1 + len(rows3), 12)
            elif k == "panel":
                sh.boxes(F_SES)
                sh.cell("A", r["time"], F_SES, font(10, True), LEFT)
                if r.get("n"): sh.cell("B", r["n"], F_SES, font(10))
                sh.cell("C", f'Moderator - {r["moderator"]}' if r.get("moderator") else "PANEL DISCUSSION", F_SES, font(10, True), LEFT)
                title = ("PANEL DISCUSSION - " if r.get("moderator") else "") + r["topic"]
                sh.cell("D", rich((title + "\n", 10, True, CRIMSON), ("Panelists - " + r["panelists"], 10, False, INK)), F_SES)
                sh.height(lines(title, WIDTHS["D"]) + lines("Panelists - " + r["panelists"], WIDTHS["D"]), 12)
            sh.next()

    ps = ws.page_setup
    ps.orientation, ps.paperSize = "landscape", ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ps.fitToWidth, ps.fitToHeight = 1, 0
    ws.print_title_rows = f"{head_row}:{head_row}"
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "IACTS TechnoCollege CME 2026  |  iactstechnocollegecme2026.com  |  Page &P of &N"
    ws.oddFooter.center.size, ws.oddFooter.center.font = 8, "Arial"
    ws.sheet_view.showGridLines = False
    ws.freeze_panes = f"A{head_row + 1}"


def breakthrough(ws):
    sh = Sheet(ws)
    widths = {"A": 12, "B": 100}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    header(sh, widths)
    for col in "AB":
        ws[f"{col}{sh.r}"].border = BOX
    ws.merge_cells(f"A{sh.r}:B{sh.r}")
    sh.cell("A", "Breakthrough Sessions - Hall B", CRIMSON, font(14, True, "FFFFFF")); ws[f"B{sh.r}"].fill = fill(CRIMSON); sh.height(1, 14); sh.next()
    head_row = sh.r
    for col, t in zip("AB", ["S. No.", "Topic Name"]):
        c = sh.cell(col, t, DEEP, font(10, True, "FFFFFF")); c.border = BOX
    sh.height(1, 10); sh.next()
    for i, t in enumerate(DATA["breakthrough"], 1):
        bg = F_ALT if i % 2 else "FFFFFF"
        for col in "AB":
            ws[f"{col}{sh.r}"].border = BOX
        sh.cell("A", i, bg, font(11)); sh.cell("B", t, bg, font(11, True)); sh.height(1, 12); sh.next()
    for col in "AB":
        ws[f"{col}{sh.r}"].border = BOX
        ws[f"{col}{sh.r}"].fill = fill(F_INFO)
    ws.merge_cells(f"A{sh.r}:B{sh.r}")
    sh.cell("A", "Breakthrough Sessions Run in Parallel with the Scientific Program", F_INFO, font(11, True, WINE)); sh.height(1, 12)

    ps = ws.page_setup
    ps.orientation, ps.paperSize = "landscape", ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ps.fitToWidth, ps.fitToHeight = 1, 1
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "IACTS TechnoCollege CME 2026  |  iactstechnocollegecme2026.com"
    ws.oddFooter.center.size, ws.oddFooter.center.font = 8, "Arial"
    ws.sheet_view.showGridLines = False


def workshop(ws):
    """The workshop day as the committee's sheet lays it out: Hall A and Hall B
    side by side (S.No · Topic · Faculty · Timings each), rows lined up."""
    sh = Sheet(ws)
    widths = {"A": 7, "B": 30, "C": 30, "D": 20, "E": 2, "F": 7, "G": 30, "H": 30, "I": 20}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    header(sh, widths)
    halls = ("A", "D"), ("F", "I")

    def band(value, f, fnt, h=16):
        for col in "ABCDEFGHI":
            c = ws[f"{col}{sh.r}"]
            c.border, c.fill = BOX, fill(f)
        ws.merge_cells(f"A{sh.r}:I{sh.r}")
        sh.cell("A", value, f, fnt); sh.height(1, h); sh.next()

    band(DATA["title"], CRIMSON, font(15, True, "FFFFFF"), 20)
    band(f'Venue : {DATA["venue"]}', F_INFO, font(11.5, True, WINE))
    band(DATA["date"], F_INFO, font(11.5, True, WINE))

    for (a, b), name in zip(halls, DATA["halls"]):                       # hall headings
        cols = "ABCDEFGHI"[ "ABCDEFGHI".index(a): "ABCDEFGHI".index(b) + 1]
        for col in cols:
            ws[f"{col}{sh.r}"].border, ws[f"{col}{sh.r}"].fill = BOX, fill(F_SES)
        ws.merge_cells(f"{a}{sh.r}:{b}{sh.r}")
        sh.cell(a, name, F_SES, font(13, True, CRIMSON))
    sh.height(1, 18); sh.next()
    head_row = sh.r
    for a, _ in halls:                                                   # column headings
        start = "ABCDEFGHI".index(a)
        for i, t in enumerate(["S.No.", "Topic", "Faculty", "Timings"]):
            c = sh.cell("ABCDEFGHI"[start + i], t, DEEP, font(10, True, "FFFFFF")); c.border = BOX
    sh.height(1, 10); sh.next()

    for i, pair in enumerate(DATA["rows"]):
        bg = F_ALT if i % 2 else "FFFFFF"
        tall = 1
        for (a, _), c in zip(halls, pair):
            s0 = "ABCDEFGHI".index(a)
            A, B, C, D = "ABCDEFGHI"[s0:s0 + 4]
            for col in (A, B, C, D):
                ws[f"{col}{sh.r}"].border = BOX
            if c is None:
                for col in (A, B, C, D):
                    ws[f"{col}{sh.r}"].fill = fill(bg)
            elif c.get("lunch"):
                for col in (A, B, C, D):
                    ws[f"{col}{sh.r}"].fill = fill(F_BRK)
                ws.merge_cells(f"{B}{sh.r}:{C}{sh.r}")
                sh.cell(B, "Lunch", F_BRK, font(11, True, WINE))
                sh.cell(D, c["time"], F_BRK, font(10, True))
            else:
                sh.cell(A, c["n"], bg, font(10))
                sh.cell(B, c["topic"], bg, font(10))
                sh.cell(C, "\n".join(c["faculty"]), bg, font(10, True), LEFT)
                sh.cell(D, c["time"], bg, font(10))
                tall = max(tall, len(c["faculty"]), lines(c["topic"], widths[B]), c["time"].count("\n") + 1)
        ws.row_dimensions[sh.r].height = tall * 19 + 30; sh.next()   # roomy: one page, filled

    ps = ws.page_setup
    ps.orientation, ps.paperSize = "landscape", ws.PAPERSIZE_A4
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ps.fitToWidth, ps.fitToHeight = 1, 1
    ws.page_margins.left = ws.page_margins.right = 0.4
    ws.page_margins.top = ws.page_margins.bottom = 0.5
    ws.print_options.horizontalCentered = True
    ws.oddFooter.center.text = "IACTS TechnoCollege CME 2026  |  iactstechnocollegecme2026.com"
    ws.oddFooter.center.size, ws.oddFooter.center.font = 8, "Arial"
    ws.sheet_view.showGridLines = False


if __name__ == "__main__":
    out = Path(sys.argv[1] if len(sys.argv) > 1 else f"{DOC}.xlsx").expanduser()
    wb = Workbook()
    ws = wb.active
    if DOC == "workshop":
        ws.title = "Workshop"
        workshop(ws)
    else:
        ws.title = "Scientific Schedule"
        schedule(ws)
        breakthrough(wb.create_sheet("Breakthrough Sessions"))
    wb.properties.title = f"{ws.title} — IACTS TechnoCollege CME 2026"
    wb.properties.creator = "IACTS TechnoCollege CME 2026"
    wb.save(out)
    print("wrote", out)
